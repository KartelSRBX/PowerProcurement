"""Text out of the files suppliers send (PDF, Excel, CSV, plain text), so pasting is not the only way in.
The text goes through the same reading step as a pasted answer, and a person still checks the result."""

import csv
import io
from pathlib import PurePath

from app.pipeline.service import PipelineError

MAX_BYTES = 5_000_000
MAX_CHARS = 19_000
MAX_PAGES = 30
MAX_SHEETS = 5
MAX_ROWS = 300
MAX_COLS = 30
KINDS = {".pdf", ".xlsx", ".xlsm", ".csv", ".txt"}


def _clip(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n[text cut: the file is longer than the reader takes]"
    return text


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1250"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise PipelineError(
                "this PDF is password protected; open it, save a copy without a password, or paste the text"
            )
        pages = [(p.extract_text() or "").strip() for p in reader.pages[:MAX_PAGES]]
    except PipelineError:
        raise
    except (PyPdfError, ValueError, KeyError, RecursionError, OSError) as exc:
        raise PipelineError("this PDF could not be read; paste the text instead") from exc
    text = "\n\n".join(p for p in pages if p)
    if not text:
        raise PipelineError(
            "this PDF has no text layer (probably a scan); paste the text or enter the numbers by hand"
        )
    return text


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.6g}" if abs(v) < 1e15 else str(v)
    return str(v).strip()


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - corrupt, password protected or not really xlsx
        raise PipelineError(
            "this Excel file could not be opened; save it as .xlsx without a password"
        ) from exc
    out: list[str] = []
    for ws in wb.worksheets[:MAX_SHEETS]:
        rows = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_ROWS:
                rows.append("[more rows not read]")
                break
            cells = [_cell(c) for c in row[:MAX_COLS]]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            out.append(f"Sheet: {ws.title}\n" + "\n".join(rows))
    wb.close()
    if not out:
        raise PipelineError(
            "this Excel file has no values to read (empty, or only formulas that were never calculated)"
        )
    return "\n\n".join(out)


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = []
    try:
        for n, row in enumerate(csv.reader(io.StringIO(text), dialect)):
            if n >= MAX_ROWS:
                rows.append("[more rows not read]")
                break
            cells = [c.strip() for c in row[:MAX_COLS]]
            if any(cells):
                rows.append(" | ".join(cells))
    except csv.Error as exc:
        raise PipelineError("this CSV could not be read; paste the text instead") from exc
    return "\n".join(rows)


def extract_text(filename: str, data: bytes) -> str:
    name = PurePath(filename or "").name[:120]
    ext = PurePath(name).suffix.lower()
    if ext == ".xls":
        raise PipelineError("old .xls files are not read; save the sheet as .xlsx and try again")
    if ext not in KINDS:
        raise PipelineError(
            "this file type is not read; use PDF, Excel (.xlsx), CSV or text, or paste the text"
        )
    if not data:
        raise PipelineError("the file is empty")
    if len(data) > MAX_BYTES:
        raise PipelineError(f"the file is larger than {MAX_BYTES // 1_000_000} MB")
    text = {".pdf": _pdf, ".xlsx": _xlsx, ".xlsm": _xlsx, ".csv": _csv}.get(ext, _decode)(data)
    text = _clip(text)
    if not text:
        raise PipelineError("no text could be read from this file")
    return f"[File: {name}]\n{text}"
