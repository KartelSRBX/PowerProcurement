"""The savings ledger. For each order a named person records a baseline, the amount the goods would have
cost without the buyer's work, and says on what basis. Code then works out two figures from data that exist:

  negotiated saving = baseline - the offer behind the order (what was agreed)
  realized saving   = baseline - the supplier's invoice, counted only once a person has accepted that invoice

A baseline is one of three kinds, so it is never a number from nowhere: the dearest offer on the same request
that covered every item, the market target price (where prices were fetched), or an amount a person types with
a note saying where it comes from. All amounts are goods without VAT, in EUR at the rates in Settings. Costs
outside the goods (freight, duty) are not in the ledger: it is a price ledger, not a landed-cost one."""

import io
from datetime import UTC, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.currency import to_eur
from app.models import Offer, PurchaseOrder, SavingsBaseline, Supplier, SupplierInvoice
from app.pipeline.invoices import compare
from app.pipeline.service import PipelineError, _audit, _median_target, _need_person, items_of
from app.pricing.service import target_total_eur

BASES = ("highest_offer", "target_price", "manual")
LABELS = {
    "highest_offer": "dearest offer that covered every item",
    "target_price": "market target price",
    "manual": "typed by a person",
}


def _order(session: Session, po_id: int) -> PurchaseOrder:
    po = session.get(PurchaseOrder, po_id)
    if po is None:
        raise LookupError("purchase order not found")
    return po


def candidates(session: Session, settings: Settings, po_id: int) -> list[dict]:
    """Baselines the data can offer for this order. Empty when it has nothing to go on."""
    po = _order(session, po_id)
    out = []
    offers = list(
        session.scalars(
            select(Offer).where(
                Offer.request_id == po.request_id, Offer.covers_all_items.is_(True), Offer.total_price > 0
            )
        )
    )
    if len(offers) >= 2:
        top = max(offers, key=lambda o: (to_eur(o.total_price, o.currency, settings), -o.id))
        sup = session.get(Supplier, top.supplier_id)
        out.append(
            {
                "basis": "highest_offer",
                "label": LABELS["highest_offer"],
                "amount_eur": round(to_eur(top.total_price, top.currency, settings), 2),
                "detail": f"{sup.name}, {top.total_price:,.2f} {top.currency}, of {len(offers)} offers",
            }
        )
    items = items_of(session, po.request_id)
    target = target_total_eur(
        session, po.request_id, {i.id: _median_target(session, i, settings) for i in items}
    )
    if target:
        out.append(
            {
                "basis": "target_price",
                "label": LABELS["target_price"],
                "amount_eur": round(target, 2),
                "detail": "median of the market prices fetched for each item, times the quantities",
            }
        )
    return out


def view(session: Session, settings: Settings, po_id: int) -> dict:
    po = _order(session, po_id)
    row = session.scalar(select(SavingsBaseline).where(SavingsBaseline.order_id == po.id))
    offer = session.get(Offer, po.offer_id)
    return {
        "order_id": po.id,
        "negotiated_eur": round(to_eur(offer.total_price, offer.currency, settings), 2),
        "baseline": None
        if row is None
        else {
            "basis": row.basis,
            "label": LABELS[row.basis],
            "amount_eur": row.amount_eur,
            "note": row.note,
            "set_by": row.set_by,
            "set_at": row.set_at.isoformat() if row.set_at else None,
        },
        "candidates": candidates(session, settings, po_id),
    }


def set_baseline(
    session: Session,
    settings: Settings,
    po_id: int,
    set_by: str,
    basis: str,
    amount_eur: float | None = None,
    note: str = "",
) -> dict:
    """Record the baseline. For a basis the data can give, the amount is taken from the data now, not from the
    caller; a typed amount needs a note saying where it comes from."""
    by = _need_person(set_by)
    po = _order(session, po_id)
    if basis not in BASES:
        raise ValueError(f"basis must be one of {', '.join(BASES)}")
    note = note.strip()[:300]
    if basis == "manual":
        if amount_eur is None or amount_eur <= 0:
            raise ValueError("a typed baseline needs an amount above zero")
        if len(note) < 3:
            raise ValueError("say where a typed baseline comes from")
        amount = round(amount_eur, 2)
    else:
        found = {c["basis"]: c for c in candidates(session, settings, po_id)}
        if basis not in found:
            raise PipelineError(f"the data has nothing to base “{LABELS[basis]}” on for this order")
        amount = found[basis]["amount_eur"]
    row = session.scalar(select(SavingsBaseline).where(SavingsBaseline.order_id == po.id))
    before = None if row is None else {"basis": row.basis, "amount_eur": row.amount_eur}
    if row is None:
        row = SavingsBaseline(order_id=po.id, basis=basis, amount_eur=amount, note=note, set_by=by)
        session.add(row)
    else:
        row.basis, row.amount_eur, row.note, row.set_by = basis, amount, note, by
        row.set_at = datetime.now(UTC)
    _audit(
        session,
        "human",
        by,
        "savings.baseline",
        order_id=po.id,
        basis=basis,
        amount_eur=amount,
        note=note,
        before=before,
    )
    session.commit()
    return view(session, settings, po_id)


def clear_baseline(session: Session, settings: Settings, po_id: int, cleared_by: str) -> dict:
    by = _need_person(cleared_by)
    po = _order(session, po_id)
    row = session.scalar(select(SavingsBaseline).where(SavingsBaseline.order_id == po.id))
    if row is None:
        raise LookupError("this order has no baseline")
    _audit(
        session,
        "human",
        by,
        "savings.baseline_cleared",
        order_id=po.id,
        basis=row.basis,
        amount_eur=row.amount_eur,
    )
    session.delete(row)
    session.commit()
    return view(session, settings, po_id)


def _pct(part: float, whole: float) -> float | None:
    return round(100 * part / whole, 1) if whole else None


def ledger(session: Session, settings: Settings) -> dict:
    """One row per order that has a baseline, and the totals. The realized figures use only invoices a person
    has accepted, and their totals only those orders, so a baseline is not set against an unseen invoice."""
    rows = []
    base_all = neg_all = 0.0
    base_real = real_all = 0.0
    real_n = differed_n = 0
    for b in session.scalars(select(SavingsBaseline).order_by(SavingsBaseline.order_id)):
        po = session.get(PurchaseOrder, b.order_id)
        offer = session.get(Offer, po.offer_id)
        sup = session.get(Supplier, po.supplier_id)
        neg = round(to_eur(offer.total_price, offer.currency, settings), 2)
        inv = session.scalar(select(SupplierInvoice).where(SupplierInvoice.order_id == po.id))
        invoiced = realized = None
        differs = False
        state = "no invoice"
        if inv is not None:
            invoiced = round(to_eur(inv.total_net, inv.currency, settings), 2)
            state = "entered, not accepted"
            if inv.accepted_by:
                state = "accepted"
                realized = round(b.amount_eur - invoiced, 2)
                differs = compare(session, settings, po, inv)["status"] == "differs"
        row = {
            "order_id": po.id,
            "request_id": po.request_id,
            "supplier": sup.name,
            "basis": b.basis,
            "basis_label": LABELS[b.basis],
            "note": b.note,
            "set_by": b.set_by,
            "baseline_eur": b.amount_eur,
            "negotiated_eur": neg,
            "negotiated_saving_eur": round(b.amount_eur - neg, 2),
            "negotiated_saving_pct": _pct(b.amount_eur - neg, b.amount_eur),
            "invoice_number": inv.number if inv else None,
            "invoice_state": state,
            "invoice_reason": inv.accept_reason if inv and inv.accepted_by else None,
            "invoice_differs": differs,
            "invoiced_eur": invoiced,
            "realized_saving_eur": realized,
            "realized_saving_pct": None if realized is None else _pct(realized, b.amount_eur),
        }
        rows.append(row)
        base_all += b.amount_eur
        neg_all += b.amount_eur - neg
        if realized is not None:
            real_n += 1
            base_real += b.amount_eur
            real_all += realized
            differed_n += 1 if differs else 0
    return {
        "rows": rows,
        "totals": {
            "orders": len(rows),
            "baseline_eur": round(base_all, 2),
            "negotiated_saving_eur": round(neg_all, 2),
            "negotiated_saving_pct": _pct(neg_all, base_all),
            "realized_orders": real_n,
            "realized_baseline_eur": round(base_real, 2),
            "realized_saving_eur": round(real_all, 2),
            "realized_saving_pct": _pct(real_all, base_real),
            "realized_accepted_with_difference": differed_n,
        },
        "orders_without_baseline": session.query(PurchaseOrder)
        .filter(PurchaseOrder.id.notin_([r["order_id"] for r in rows] or [0]))
        .count(),
        "note": "Goods without VAT, in EUR at the rates in Settings. Realized counts accepted invoices only.",
    }


def ledger_xlsx(session: Session, settings: Settings) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    data = ledger(session, settings)
    wb = Workbook()
    ws = wb.active
    ws.title = "Savings"
    head = [
        "Order",
        "Supplier",
        "Baseline basis",
        "Baseline note",
        "Baseline set by",
        "Baseline EUR",
        "Negotiated EUR",
        "Negotiated saving EUR",
        "Negotiated saving %",
        "Invoice",
        "Invoice state",
        "Invoiced EUR",
        "Realized saving EUR",
        "Realized saving %",
    ]
    ws.append(head)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
    for r in data["rows"]:
        ws.append(
            [
                f"PO #{r['order_id']}",
                r["supplier"],
                r["basis_label"],
                r["note"],
                r["set_by"],
                r["baseline_eur"],
                r["negotiated_eur"],
                r["negotiated_saving_eur"],
                r["negotiated_saving_pct"],
                r["invoice_number"] or "",
                r["invoice_state"],
                r["invoiced_eur"] if r["invoiced_eur"] is not None else "",
                r["realized_saving_eur"] if r["realized_saving_eur"] is not None else "",
                r["realized_saving_pct"] if r["realized_saving_pct"] is not None else "",
            ]
        )
    t = data["totals"]
    ws.append([])
    ws.append(
        [
            "Total",
            f"{t['orders']} order(s)",
            "",
            "",
            "",
            t["baseline_eur"],
            "",
            t["negotiated_saving_eur"],
            t["negotiated_saving_pct"],
        ]
    )
    ws.append(
        [
            "Realized",
            f"{t['realized_orders']} order(s) with an accepted invoice",
            "",
            "",
            "",
            t["realized_baseline_eur"],
            "",
            "",
            "",
            "",
            "",
            "",
            t["realized_saving_eur"],
            t["realized_saving_pct"],
        ]
    )
    for cell in ws[ws.max_row - 1] + ws[ws.max_row]:
        cell.font = bold
    ws.append([])
    ws.append([data["note"]])
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
