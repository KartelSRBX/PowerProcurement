"""Offers side by side, item by item. Code does the arithmetic and the "can these be compared" check;
a person reads the table. Prices are brought to EUR at the fixed rates in Settings, and an offer that
cannot honestly be set against the others says why instead of sneaking into the "cheapest" column."""

import io
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.currency import to_eur
from app.models import Offer, OfferLine, Supplier
from app.pipeline.service import get_request, items_of, term_flags

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TOTAL_TOLERANCE = 0.01  # a stated total may differ from the sum of the lines by 1% (rounding)


def compare_offers(session: Session, settings: Settings, request_id: int, today: date | None = None) -> dict:
    get_request(session, request_id)
    today = today or date.today()
    items = items_of(session, request_id)
    offers = list(session.scalars(select(Offer).where(Offer.request_id == request_id).order_by(Offer.id)))
    lines: dict[int, dict[int, OfferLine]] = {o.id: {} for o in offers}
    for ln in session.scalars(select(OfferLine).where(OfferLine.offer_id.in_(list(lines) or [0]))):
        lines[ln.offer_id][ln.item_id] = ln
    suppliers = {o.id: session.get(Supplier, o.supplier_id) for o in offers}
    rules = {o.incoterm for o in offers if o.incoterm}

    cells: dict[int, dict[int, dict]] = {}
    rows = []
    for o in offers:
        sup = suppliers[o.id]
        own = lines[o.id]
        cells[o.id] = {}
        lines_total = 0.0
        for it in items:
            ln = own.get(it.id)
            if ln is None:
                continue
            unit_eur = to_eur(ln.unit_price, o.currency, settings)
            lines_total += ln.unit_price * (it.quantity or 0)
            cells[o.id][it.id] = {
                "unit_price": ln.unit_price,
                "unit_price_eur": round(unit_eur, 4),
                "line_total_eur": round(unit_eur * (it.quantity or 0), 2),
                "moq": ln.moq,
                "moq_ok": ln.moq is None or (it.quantity or 0) >= ln.moq,
                "best": False,
            }
        blockers, notes = [], []
        if not own:
            blockers.append("only a total price: no price per item to compare")
        else:
            missing = [it.description for it in items if it.id not in own]
            if missing:
                blockers.append(f"no price for: {', '.join(missing)}")
            if abs(lines_total - o.total_price) > TOTAL_TOLERANCE * o.total_price:
                notes.append(f"stated total {o.total_price:g} differs from the lines' sum {lines_total:.2f}")
            for it in items:
                c = cells[o.id].get(it.id)
                if c and not c["moq_ok"]:
                    notes.append(
                        f"minimum order {c['moq']:g} is above the {it.quantity:g} asked for: {it.description}"
                    )
        if len(rules) > 1 and o.incoterm:
            others = ", ".join(sorted(rules - {o.incoterm}))
            blockers.append(f"{o.incoterm} here but {others} in other offers: freight differs")
        elif rules and not o.incoterm:
            blockers.append("no Incoterms rule while other offers state one: freight may be missing")
        if o.valid_until and o.valid_until < today:
            blockers.append(f"price expired on {o.valid_until.isoformat()}")
        if sup.sanctioned:
            blockers.append("supplier is on the sanctions list")
        rows.append(
            {
                "offer_id": o.id,
                "supplier": sup.name,
                "country": sup.country,
                "currency": o.currency,
                "total": o.total_price,
                "total_eur": round(to_eur(o.total_price, o.currency, settings), 2),
                "lead_time_days": o.lead_time_days,
                "incoterm": o.incoterm,
                "incoterm_place": o.incoterm_place,
                "payment_days": o.payment_days,
                "valid_until": o.valid_until,
                "has_lines": bool(own),
                "comparable": not blockers,
                "not_comparable": blockers,
                "notes": notes,
                "term_flags": term_flags(o, today),
            }
        )

    comparable = {r["offer_id"] for r in rows if r["comparable"]}
    best: dict[int, int] = {}
    for it in items:
        priced = [(cells[oid][it.id]["unit_price_eur"], oid) for oid in comparable if it.id in cells[oid]]
        if len(priced) >= 2:
            low, oid = min(priced)
            if sum(1 for p, _ in priced if p == low) == 1:  # a tie has no single cheapest
                cells[oid][it.id]["best"] = True
                best[it.id] = oid
    return {
        "request_id": request_id,
        "currency": "EUR",
        "items": [
            {
                "id": i.id,
                "description": i.description,
                "part_number": i.part_number,
                "quantity": i.quantity,
                "unit": i.unit,
            }
            for i in items
        ],
        "offers": rows,
        "cells": {str(oid): {str(iid): c for iid, c in cs.items()} for oid, cs in cells.items()},
        "best_by_item": {str(k): v for k, v in best.items()},
    }


def comparison_xlsx(session: Session, settings: Settings, request_id: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    cmp = compare_offers(session, settings, request_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    bold, green = Font(bold=True), PatternFill("solid", fgColor="C6EFCE")
    ws.append(
        ["Item", "Part number", "Quantity", "Unit"]
        + [f"{o['supplier']} ({o['currency']}→EUR)" for o in cmp["offers"]]
    )
    for c in ws[1]:
        c.font = bold
    for it in cmp["items"]:
        row = [it["description"], it["part_number"] or "", it["quantity"], it["unit"] or ""]
        for o in cmp["offers"]:
            cell = cmp["cells"][str(o["offer_id"])].get(str(it["id"]))
            row.append(cell["unit_price_eur"] if cell else "no price")
        ws.append(row)
        for col, o in enumerate(cmp["offers"], start=5):
            cell = cmp["cells"][str(o["offer_id"])].get(str(it["id"]))
            if cell and cell["best"]:
                ws.cell(row=ws.max_row, column=col).fill = green
    ws.append([])
    extra = [
        ("Total in the offer's currency", lambda o: f"{o['total']:g} {o['currency']}"),
        ("Total, EUR", lambda o: o["total_eur"]),
        ("Lead time, days", lambda o: o["lead_time_days"]),
        ("Incoterms", lambda o: f"{o['incoterm'] or 'not stated'} {o['incoterm_place'] or ''}".strip()),
        ("Payment, days", lambda o: o["payment_days"] if o["payment_days"] is not None else "not stated"),
        ("Valid until", lambda o: o["valid_until"].isoformat() if o["valid_until"] else "not stated"),
        ("Comparable", lambda o: "yes" if o["comparable"] else "NO"),
        ("Why not comparable", lambda o: "; ".join(o["not_comparable"])),
        ("Notes", lambda o: "; ".join(o["notes"])),
    ]
    for label, fn in extra:
        ws.append([label, "", "", ""] + [fn(o) for o in cmp["offers"]])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    ws.append(["Unit prices in EUR at the Settings rates; green is the lowest among comparable offers."])
    ws.column_dimensions["A"].width = 36
    for col in range(2, 5 + len(cmp["offers"])):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
