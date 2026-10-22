"""Prices per item, the side-by-side comparison and its "not comparable" reasons."""

import io
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.test_pipeline_e2e import SETTINGS, client, new_request  # noqa: F401


def _setup(c):
    rid = new_request(c)
    c.post(f"/requests/{rid}/discover")
    for d in c.post(f"/requests/{rid}/rfqs/draft").json():
        c.post(f"/rfqs/{d['id']}/decision", json={"approve": True, "decided_by": "kartel"})
    suppliers = [s["supplier_id"] for s in c.get(f"/requests/{rid}/shortlist").json()]
    return rid, suppliers, c.get(f"/requests/{rid}/items").json()


def _offer(c, rid, sid, prices, **over):
    """prices: unit price per item in the order of the request's items."""
    items = c.get(f"/requests/{rid}/items").json()
    lines = [{"item_id": i["id"], "unit_price": p} for i, p in zip(items, prices, strict=True) if p]
    total = sum(p * (i["quantity"] or 0) for i, p in zip(items, prices, strict=True) if p)
    body = {"supplier_id": sid, "total_price": total, "currency": "EUR", "lead_time_days": 10, "lines": lines}
    r = c.post(f"/requests/{rid}/offers", json={**body, **over})
    assert r.status_code == 200, r.text
    return r.json()


def test_lines_decide_coverage_and_the_cheapest_is_marked_per_item(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    a = _offer(c, rid, sids[0], [10, 20], incoterm="DAP", incoterm_place="Belgrade")
    b = _offer(c, rid, sids[1], [12, 15], incoterm="DAP", incoterm_place="Belgrade")
    assert a["covers_all_items"] and b["covers_all_items"]
    cmp = c.get(f"/requests/{rid}/comparison").json()
    assert [o["comparable"] for o in cmp["offers"]] == [True, True]
    first, second = (str(i["id"]) for i in items)
    assert cmp["best_by_item"][first] == a["id"] and cmp["best_by_item"][second] == b["id"]
    assert cmp["cells"][str(a["id"])][first]["best"] and not cmp["cells"][str(b["id"])][first]["best"]


def test_missing_item_overrides_the_ticked_covers_all_items(client):  # noqa: F811
    c = client
    rid, sids, _ = _setup(c)
    partial = _offer(c, rid, sids[0], [10, 0], covers_all_items=True)
    assert partial["covers_all_items"] is False  # code decides, not the tick box
    cmp = c.get(f"/requests/{rid}/comparison").json()
    assert not cmp["offers"][0]["comparable"] and "no price for" in cmp["offers"][0]["not_comparable"][0]


def test_total_only_offer_is_not_comparable_and_never_cheapest(client):  # noqa: F811
    c = client
    rid, sids, _ = _setup(c)
    _offer(c, rid, sids[0], [10, 20])
    c.post(
        f"/requests/{rid}/offers",
        json={"supplier_id": sids[1], "total_price": 1, "currency": "EUR", "lead_time_days": 3},
    )
    cmp = c.get(f"/requests/{rid}/comparison").json()
    bare = cmp["offers"][1]
    assert bare["comparable"] is False and "only a total" in bare["not_comparable"][0]
    assert cmp["best_by_item"] == {}  # one comparable offer is nothing to compare against


def test_different_incoterms_and_expired_prices_block_comparison(client):  # noqa: F811
    c = client
    rid, sids, _ = _setup(c)
    old = (date.today() - timedelta(days=1)).isoformat()
    _offer(c, rid, sids[0], [10, 20], incoterm="EXW", incoterm_place="Novi Sad")
    _offer(c, rid, sids[1], [9, 18], incoterm="DAP", incoterm_place="Belgrade", valid_until=old)
    why = [" | ".join(o["not_comparable"]) for o in c.get(f"/requests/{rid}/comparison").json()["offers"]]
    assert "EXW here but DAP" in why[0] and "DAP here but EXW" in why[1] and "expired" in why[1]


def test_currencies_are_aligned_to_eur_and_a_wrong_total_is_noted(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    usd = _offer(c, rid, sids[0], [10, 20], currency="USD", total_price=1)
    cmp = c.get(f"/requests/{rid}/comparison").json()
    cell = cmp["cells"][str(usd["id"])][str(items[0]["id"])]
    assert cell["unit_price"] == 10 and cell["unit_price_eur"] == round(10 / SETTINGS.fx_usd_per_eur, 4)
    assert any("differs from the lines" in n for n in cmp["offers"][0]["notes"])


def test_moq_above_quantity_makes_moq_ok_false(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    lines = [{"item_id": i["id"], "unit_price": 5, "moq": (i["quantity"] or 0) + 1} for i in items]
    r = c.post(
        f"/requests/{rid}/offers",
        json={
            "supplier_id": sids[0],
            "total_price": 100,
            "currency": "EUR",
            "lead_time_days": 5,
            "moq_ok": True,
            "lines": lines,
        },
    ).json()
    assert r["moq_ok"] is False


def test_bad_lines_are_refused(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    base = {"supplier_id": sids[0], "total_price": 10, "currency": "EUR", "lead_time_days": 5}
    i0 = items[0]["id"]
    for lines in (
        [{"item_id": 999999, "unit_price": 1}],
        [{"item_id": i0, "unit_price": 1}, {"item_id": i0, "unit_price": 2}],
        [{"item_id": i0, "unit_price": 0}],
        [{"item_id": i0, "unit_price": 1, "moq": -1}],
    ):
        assert c.post(f"/requests/{rid}/offers", json={**base, "lines": lines}).status_code == 409, lines
    assert c.get(f"/requests/{rid}/offers").json() == []


def test_confirming_a_reply_stores_and_replaces_lines(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    r = c.post(
        f"/requests/{rid}/replies",
        json={"supplier_id": sids[0], "text": "see prices", "entered_by": "kartel"},
    ).json()
    ln = [{"item_id": i["id"], "unit_price": 3} for i in items]
    o = c.post(
        f"/replies/{r['id']}/confirm",
        json={
            "confirmed_by": "kartel",
            "total_price": 99,
            "currency": "EUR",
            "lead_time_days": 4,
            "lines": ln,
        },
    ).json()
    r2 = c.post(
        f"/requests/{rid}/replies",
        json={"supplier_id": sids[0], "text": "new prices", "entered_by": "kartel"},
    ).json()
    ln2 = [{"item_id": i["id"], "unit_price": 2} for i in items]
    c.post(
        f"/replies/{r2['id']}/confirm",
        json={
            "confirmed_by": "kartel",
            "total_price": 50,
            "currency": "EUR",
            "lead_time_days": 4,
            "lines": ln2,
        },
    )
    cells = c.get(f"/requests/{rid}/comparison").json()["cells"][str(o["id"])]
    assert {v["unit_price"] for v in cells.values()} == {2}


def test_excel_export_has_items_offers_and_reasons(client):  # noqa: F811
    c = client
    rid, sids, items = _setup(c)
    _offer(c, rid, sids[0], [10, 20])
    _offer(c, rid, sids[1], [12, 15])
    r = c.get(f"/requests/{rid}/comparison.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    labels = [row[0].value for row in ws.iter_rows()]
    assert "Comparable" in labels and "Incoterms" in labels
    assert ws.cell(row=2, column=1).value == items[0]["description"]
    assert ws.cell(row=2, column=5).value == 10
