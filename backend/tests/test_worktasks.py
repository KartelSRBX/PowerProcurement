"""Tasks in Bitrix24 (scripted portal), the demo stand-in and the Excel mode. No internet is used."""

# ruff: noqa: E501
import base64
import io
import json
from urllib.parse import parse_qs

import httpx
import pytest
from cryptography.fernet import Fernet
from openpyxl import Workbook, load_workbook

from app.config import get_settings
from app.main import app
from app.models import AppSetting
from app.worktasks import api as tasks_api
from app.worktasks.providers import BitrixProvider
from tests.test_pipeline_e2e import SETTINGS, client, new_request  # noqa: F401

HOOK = "https://acme.bitrix24.com/rest/7/secrettoken/"


class Portal:
    """Just enough of Bitrix24's REST answers. `status` is per task id, so a test can play the colleague."""

    def __init__(self):
        self.tasks: dict[str, dict] = {}
        self.calls: list[tuple[str, dict]] = []
        self.fail_on: str | None = None
        self.throttle = 0

    def __call__(self, req: httpx.Request) -> httpx.Response:
        method = req.url.path.rsplit("/", 1)[-1].removesuffix(".json")
        body = json.loads(req.content or b"{}")
        self.calls.append((method, body))
        if self.throttle and method == "profile":
            self.throttle -= 1
            return httpx.Response(503, json={"error": "QUERY_LIMIT_EXCEEDED", "error_description": "slow"})
        if self.fail_on == method:
            return httpx.Response(401, json={"error": "insufficient_scope", "error_description": "x"})
        if method == "profile":
            return httpx.Response(200, json={"result": {"ID": "7", "NAME": "Ana", "LAST_NAME": "Petrovic"}})
        if method == "tasks.task.add":
            f = body["fields"]
            tid = str(100 + len(self.tasks))
            self.tasks[tid] = {
                "id": tid,
                "title": f["TITLE"],
                "status": "2",
                "responsibleId": f.get("RESPONSIBLE_ID", "7"),
                "createdBy": f.get("CREATED_BY", "7"),
                "deadline": f.get("DEADLINE"),
                "stageId": "0",
                "fields": f,
            }
            return httpx.Response(200, json={"result": {"task": {"id": tid}}})
        if method == "task.checklistitem.add":
            return httpx.Response(200, json={"result": 1})
        if method == "batch":
            out = {}
            for name, cmd in body["cmd"].items():
                tid = parse_qs(cmd.split("?", 1)[1])["taskId"][0]
                if tid in self.tasks:
                    out[name] = {"task": {k: v for k, v in self.tasks[tid].items() if k != "fields"}}
            return httpx.Response(200, json={"result": {"result": out, "result_error": {}}})
        if method.startswith("tasks.task."):
            tid = str(body["taskId"])
            code = {"start": "3", "complete": "5", "pause": "6", "renew": "2"}[method.rsplit(".", 1)[1]]
            self.tasks[tid]["status"] = code
            return httpx.Response(200, json={"result": {"task": {"id": tid}}})
        if method == "user.get":
            return httpx.Response(
                200, json={"result": [{"ID": "7", "NAME": "Ana", "LAST_NAME": "P", "WORK_POSITION": "CEO"}]}
            )
        if method == "sonet_group.get":
            return httpx.Response(200, json={"result": [{"ID": "12", "NAME": "Sales"}]})
        if method == "task.stages.get":
            return httpx.Response(200, json={"result": {"31": {"ID": "31", "TITLE": "In work"}}})
        return httpx.Response(404, json={"error": "ERROR_METHOD_NOT_FOUND", "error_description": "no"})


@pytest.fixture
def portal(client, monkeypatch):  # noqa: F811
    monkeypatch.setattr(BitrixProvider, "PAUSE", 0)
    monkeypatch.setattr("app.worktasks.providers.time.sleep", lambda s: None)
    p = Portal()
    monkeypatch.setattr(tasks_api, "TRANSPORT", httpx.MockTransport(p))
    key = Fernet.generate_key().decode()
    app.dependency_overrides[get_settings] = lambda: SETTINGS.model_copy(update={"app_secret_key": key})
    return client, p


def configure(c, **over):
    cur = c.get("/tasks/config").json()["config"]
    cur.update({"provider": "bitrix", **over})
    r = c.put("/tasks/config", json={"changed_by": "kartel", "config": cur})
    assert r.status_code == 200, r.text
    return r.json()["config"]


def connect(c):
    r = c.put("/tasks/webhook", json={"changed_by": "kartel", "url": HOOK})
    assert r.status_code == 200 and r.json()["webhook"]["set"] is True
    assert "secrettoken" not in json.dumps(r.json())  # the address is never shown again
    row = c.session.get(AppSetting, "tasks_bitrix_webhook")
    assert row.secret and "secrettoken" not in row.value  # and stored encrypted


def test_default_plan_is_a_tree_with_the_requests_own_words(client):  # noqa: F811
    rid = new_request(client)
    plan = client.get(f"/requests/{rid}/tasks/plan").json()
    assert [p["key"] for p in plan][0] == "request" and len(plan) == 7
    assert plan[0]["title"].startswith(f"PP-{rid:05d} · RFQ")
    assert "500 × STM32F407VGT6" in plan[0]["description"] and all(p["exists"] is False for p in plan)
    assert {p["parent"] for p in plan[1:]} == {"request"}


def test_config_is_validated_and_needs_a_name(client):  # noqa: F811
    cfg = client.get("/tasks/config").json()["config"]
    bad = dict(cfg, templates=cfg["templates"] + [dict(cfg["templates"][1], key="suppliers")])
    assert client.put("/tasks/config", json={"changed_by": "k", "config": bad}).status_code == 422
    loop = json.loads(json.dumps(cfg))
    loop["templates"][0]["parent"] = "quote"
    assert "ancestor" in client.put("/tasks/config", json={"changed_by": "k", "config": loop}).text
    ghost = json.loads(json.dumps(cfg))
    ghost["templates"][1]["role"] = "nobody"
    assert client.put("/tasks/config", json={"changed_by": "k", "config": ghost}).status_code == 422
    badmap = dict(cfg, status_map={"3": "flying"})
    assert client.put("/tasks/config", json={"changed_by": "k", "config": badmap}).status_code == 422
    assert client.put("/tasks/config", json={"changed_by": " ", "config": cfg}).status_code == 422


def test_a_switched_off_step_hands_its_children_to_the_grandparent(client):  # noqa: F811
    rid = new_request(client)
    cfg = client.get("/tasks/config").json()["config"]
    for t in cfg["templates"]:
        if t["key"] == "request":
            t["enabled"] = False
    cfg["provider"] = "excel"
    client.put("/tasks/config", json={"changed_by": "kartel", "config": cfg})
    plan = client.get(f"/requests/{rid}/tasks/plan").json()
    assert "request" not in [p["key"] for p in plan] and all(p["parent"] is None for p in plan)


def test_bitrix_creates_the_tree_with_parents_people_and_fields(portal):
    c, p = portal
    rid = new_request(c)
    cfg = c.get("/tasks/config").json()["config"]
    roles = {
        k: dict(v, user_id=u) for (k, v), u in zip(cfg["roles"].items(), ("21", "22", "23"), strict=True)
    }
    configure(
        c,
        roles=roles,
        group_id="12",
        creator_user_id="7",
        extra_fields={"PRIORITY": "2"},
        custom_fields={"ref": "UF_PP_REF"},
    )
    assert c.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).status_code == 409  # no webhook yet
    connect(c)
    assert "Ana Petrovic" in c.post("/tasks/test").json()["detail"]
    made = c.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).json()
    assert len(made["created"]) == 7 and made["created"][0] == "request"
    adds = [b["fields"] for m, b in p.calls if m == "tasks.task.add"]
    root, first = adds[0], adds[1]
    assert "PARENT_ID" not in root and first["PARENT_ID"] == "100"  # subtasks hang under the first task
    assert first["GROUP_ID"] == "12" and first["CREATED_BY"] == "7" and first["PRIORITY"] == "2"
    assert first["UF_PP_REF"] == f"PP-{rid:05d}" and first["RESPONSIBLE_ID"] == "22"  # buyer
    assert "approved by kartel" in first["DESCRIPTION"] and first["DEADLINE"]
    assert sum(m == "task.checklistitem.add" for m, _ in p.calls) == 4
    assert c.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).json()["created"] == []  # nothing twice
    assert c.post(f"/requests/{rid}/tasks", json={"by": " "}).status_code == 409


def test_a_failure_midway_keeps_what_was_made_and_continues_later(portal):
    c, p = portal
    rid = new_request(c)
    configure(c, default_responsible="7")
    connect(c)
    orig = p.__call__
    n = {"adds": 0}

    def flaky(req):
        if req.url.path.endswith("tasks.task.add.json"):
            n["adds"] += 1
            if n["adds"] == 4:
                return httpx.Response(401, json={"error": "ACCESS_DENIED", "error_description": "no"})
        return orig(req)

    tasks_api.TRANSPORT = httpx.MockTransport(flaky)
    r = c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    assert r.status_code == 409 and "3 of 7" in r.json()["detail"]
    assert c.get(f"/requests/{rid}/tasks").json()["count"] == 3
    tasks_api.TRANSPORT = httpx.MockTransport(p)
    assert len(c.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).json()["created"]) == 4
    assert c.get(f"/requests/{rid}/tasks").json()["count"] == 7


def test_sync_reads_status_people_and_deadline_and_keeps_the_history(portal):
    c, p = portal
    rid = new_request(c)
    configure(c, default_responsible="7")
    connect(c)
    c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    t = c.get(f"/requests/{rid}/tasks").json()
    kids = t["tasks"][0]["children"]
    assert t["progress"] == 0 and len(kids) == 6
    ext = {k["key"]: k["external_id"] for k in kids}
    p.tasks[ext["suppliers"]].update(status="5", responsibleId="30")
    p.tasks[ext["offers"]]["status"] = "3"
    p.tasks[ext["quote"]]["status"] = "99"  # a status this company invented
    r = c.post(f"/requests/{rid}/tasks/sync").json()
    assert r["checked"] == 7 and r["changed"] == 2 and r["unmapped_statuses"] == ["99"]
    t = c.get(f"/requests/{rid}/tasks").json()
    by = {k["key"]: k for k in t["tasks"][0]["children"]}
    assert by["suppliers"]["state"] == "done" and by["suppliers"]["responsible_id"] == "30"
    kinds = [(e["kind"], e["new"]) for e in by["suppliers"]["events"]]
    assert ("state", "done") in kinds and ("responsible", "30") in kinds
    assert by["quote"]["state"] == "todo"  # unknown status changes nothing but is reported
    assert t["progress"] == round((100 + 50) / 6)
    # a company with its own meaning for 99 adds it to the map and the next sync follows
    cfg = c.get("/tasks/config").json()["config"]
    cfg["status_map"]["99"] = "review"
    c.put("/tasks/config", json={"changed_by": "kartel", "config": cfg})
    assert c.post(f"/requests/{rid}/tasks/sync").json()["unmapped_statuses"] == []
    by = {k["key"]: k for k in c.get(f"/requests/{rid}/tasks").json()["tasks"][0]["children"]}
    assert by["quote"]["state"] == "review"
    del p.tasks[ext["invoice"]]  # deleted in Bitrix24
    assert c.post(f"/requests/{rid}/tasks/sync").json()["missing"] == ["invoice"]


def test_kanban_stage_wins_over_status_when_the_company_maps_stages(portal):
    c, p = portal
    rid = new_request(c)
    configure(c, stage_map={"31": "doing"}, default_responsible="7")
    connect(c)
    c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    tid = next(iter(p.tasks))
    p.tasks[tid]["stageId"] = "31"
    c.post(f"/requests/{rid}/tasks/sync")
    assert c.get(f"/requests/{rid}/tasks").json()["tasks"][0]["state"] == "doing"


def test_changing_a_state_goes_through_the_remote_system_and_needs_a_name(portal):
    c, p = portal
    rid = new_request(c)
    configure(c, default_responsible="7")
    connect(c)
    c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    task = c.get(f"/requests/{rid}/tasks").json()["tasks"][0]
    assert c.post(f"/tasks/{task['id']}/state", json={"state": "done", "by": " "}).status_code == 409
    assert c.post(f"/tasks/{task['id']}/state", json={"state": "review", "by": "kartel"}).status_code == 409
    assert (
        c.post(f"/tasks/{task['id']}/state", json={"state": "done", "by": "kartel"}).json()["state"] == "done"
    )
    assert p.tasks[task["external_id"]]["status"] == "5" and p.calls[-1][0] == "tasks.task.complete"


def test_drift_shows_when_the_work_is_ahead_or_behind_the_app(portal):
    c, p = portal
    rid = new_request(c)
    configure(c, default_responsible="7")
    connect(c)
    c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    by = {k["key"]: k for k in c.get(f"/requests/{rid}/tasks").json()["tasks"][0]["children"]}
    assert all(k["drift"] is None for k in by.values())  # the request has only just been confirmed
    c.post(f"/requests/{rid}/tasks/sync")
    ext = by["offers"]["external_id"]
    p.tasks[ext]["status"] = "5"  # finished in Bitrix24 although no offer exists in the app
    c.post(f"/requests/{rid}/tasks/sync")
    by = {k["key"]: k for k in c.get(f"/requests/{rid}/tasks").json()["tasks"][0]["children"]}
    assert by["offers"]["drift"] == "ahead"


def test_bitrix_errors_are_explained_and_throttling_is_retried(portal):
    c, p = portal
    configure(c)
    connect(c)
    p.throttle = 2
    assert c.post("/tasks/test").json()["ok"] is True
    p.fail_on = "profile"
    r = c.post("/tasks/test")
    assert r.status_code == 409 and "tick" in r.json()["detail"]
    d = c.get("/tasks/directory")
    assert (
        d.status_code == 200
        and d.json()["users"][0]["name"] == "Ana P"
        and d.json()["groups"][0]["id"] == "12"
    )
    assert c.put("/tasks/webhook", json={"changed_by": "k", "url": "http://x/rest/1/t/"}).status_code == 422
    assert c.put("/tasks/webhook", json={"changed_by": "k", "url": None}).json()["webhook"]["set"] is False


def test_demo_mode_shows_the_whole_life_of_a_task_without_a_portal(client):  # noqa: F811
    rid = new_request(client)
    cfg = client.get("/tasks/config").json()["config"]
    cfg["provider"] = "demo"
    client.put("/tasks/config", json={"changed_by": "kartel", "config": cfg})
    assert client.post("/tasks/test").json()["ok"] is True
    assert len(client.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).json()["created"]) == 7
    for _ in range(5):
        client.post(f"/requests/{rid}/tasks/demo-advance")
        client.post(f"/requests/{rid}/tasks/sync")
    t = client.get(f"/requests/{rid}/tasks").json()
    assert t["progress"] > 50
    first = t["tasks"][0]["children"][0]
    assert [e["new"] for e in first["events"] if e["kind"] == "state"] == ["doing", "review", "done"]
    ov = client.get("/tasks").json()
    assert ov["requests"][0]["ref"] == f"PP-{rid:05d}" and ov["tasks"] == 7


def book(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def test_excel_mode_exports_and_reads_statuses_back_from_our_file_and_a_bitrix_export(client):  # noqa: F811
    rid = new_request(client)
    cfg = client.get("/tasks/config").json()["config"]
    cfg["provider"] = "excel"
    client.put("/tasks/config", json={"changed_by": "kartel", "config": cfg})
    assert len(client.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).json()["created"]) == 7
    assert client.post(f"/requests/{rid}/tasks/sync").status_code == 409  # nothing to fetch in file mode
    out = client.get(f"/requests/{rid}/tasks/export.xlsx")
    assert out.status_code == 200 and out.headers["content-type"].startswith("application/vnd")
    ws = load_workbook(io.BytesIO(out.content)).active
    rows = list(ws.values)
    assert rows[0][:3] == ("Request", "Key", "Task") and len(rows) == 8
    ref = f"PP-{rid:05d}"
    edited = book(
        [
            list(rows[0]),
            [ref, "suppliers", "x", "", "", "", "", "", "done"],
            [ref, "offers", "x", "", "", "", "", "", "in progress"],
            [ref, "quote", "x", "", "", "", "", "", "???"],
            ["PP-99999", "quote", "x", "", "", "", "", "", "done"],
        ]
    )
    r = client.post(
        "/tasks/import", json={"filename": "t.xlsx", "content_base64": edited, "by": "kartel"}
    ).json()
    assert r["updated"] == 2 and r["unknown_statuses"] == ["???"] and r["not_found"] == 1
    t = client.get(f"/requests/{rid}/tasks").json()
    assert {k["key"]: k["state"] for k in t["tasks"][0]["children"]}["suppliers"] == "done"
    # Bitrix24's own export: numbers instead of our keys, statuses in Russian; matched on a task number
    one = client.session.get(__import__("app.models", fromlist=["ExternalTask"]).ExternalTask, 3)
    one.external_id = "4711"
    client.session.commit()
    bitrix = book([["ID", "Название", "Статус"], ["4711", "x", "Завершена"], ["#9999", "y", "Новая"]])
    r = client.post(
        "/tasks/import", json={"filename": "b.xlsx", "content_base64": bitrix, "by": "kartel"}
    ).json()
    assert r["updated"] == 1 and r["not_found"] == 1
    bad = client.post(
        "/tasks/import", json={"filename": "x.xlsx", "content_base64": book([["a", "b"]]), "by": "k"}
    )
    assert bad.status_code == 409 and "status column" in bad.json()["detail"]
    assert (
        client.post("/tasks/import", json={"filename": "x.csv", "content_base64": "", "by": "k"}).status_code
        == 409
    )
    assert (
        client.post(
            "/tasks/import", json={"filename": "x.xlsx", "content_base64": "!!", "by": "k"}
        ).status_code
        == 409
    )


def test_bitrix_needs_somebody_responsible_before_anything_is_created(portal):
    c, p = portal
    rid = new_request(c)
    configure(c)
    connect(c)
    r = c.post(f"/requests/{rid}/tasks", json={"by": "kartel"})
    assert r.status_code == 409 and "nobody responsible" in r.json()["detail"]
    assert not any(m == "tasks.task.add" for m, _ in p.calls)


def test_off_by_default_and_a_request_must_be_confirmed(client):  # noqa: F811
    rid = new_request(client)
    assert client.get("/tasks/config").json()["config"]["provider"] == "off"
    assert client.post(f"/requests/{rid}/tasks", json={"by": "kartel"}).status_code == 409
