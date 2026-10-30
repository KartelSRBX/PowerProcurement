"""The savings ledger: a baseline a named person sets (from the data or typed with a note), the negotiated
saving (baseline minus the offer behind the order) and the realized saving (baseline minus an accepted
invoice)."""

import io

from openpyxl import load_workbook

from app.models import AuditLog
from tests.test_comparison import _offer
from tests.test_delivery import _po
from tests.test_invoice import _inv
from tests.test_pipeline_e2e import client  # noqa: F401


def _base(c, po, **over):
    body = {
        "set_by": "kartel",
        "basis": "manual",
        "amount_eur": 50000,
        "note": "last year's price list",
        **over,
    }
    return c.put(f"/orders/{po['id']}/savings/baseline", json=body)


def _accept(c, po, reason=""):
    return c.post(f"/orders/{po['id']}/supplier-invoice/accept", json={"by": "kartel", "reason": reason})


def test_an_order_without_a_baseline_is_not_in_the_ledger(client):  # noqa: F811
    c = client
    po = _po(c)
    v = c.get(f"/orders/{po['id']}/savings").json()
    assert v["baseline"] is None and v["negotiated_eur"] == 45000
    led = c.get("/savings").json()
    assert led["rows"] == [] and led["orders_without_baseline"] == 1 and led["totals"]["orders"] == 0


def test_the_ledger_matches_a_hand_calculation(client):  # noqa: F811
    c = client
    a, b = _po(c), _po(c)
    assert _base(c, a).status_code == 200  # 50,000 against 45,000 agreed
    assert _base(c, b, amount_eur=60000, note="supplier's list price").status_code == 200
    _inv(c, a, number="A-1", total_net=45000)
    _inv(c, b, number="B-1", total_net=47000)
    # nothing is realized until a person accepts the invoice
    led = c.get("/savings").json()
    assert [r["realized_saving_eur"] for r in led["rows"]] == [None, None]
    assert [r["invoice_state"] for r in led["rows"]] == ["entered, not accepted"] * 2
    assert _accept(c, a).status_code == 200
    assert _accept(c, b, "freight agreed").status_code == 200
    led = c.get("/savings").json()
    ra, rb = led["rows"]
    assert ra["negotiated_saving_eur"] == 5000 and ra["negotiated_saving_pct"] == 10.0
    assert (
        ra["realized_saving_eur"] == 5000 and ra["realized_saving_pct"] == 10.0 and not ra["invoice_differs"]
    )
    assert rb["negotiated_saving_eur"] == 15000 and rb["negotiated_saving_pct"] == 25.0
    assert rb["realized_saving_eur"] == 13000 and rb["invoice_differs"] is True
    t = led["totals"]
    assert t["orders"] == 2 and t["baseline_eur"] == 110000
    assert t["negotiated_saving_eur"] == 20000 and t["negotiated_saving_pct"] == 18.2
    assert t["realized_orders"] == 2 and t["realized_baseline_eur"] == 110000
    assert t["realized_saving_eur"] == 18000 and t["realized_saving_pct"] == 16.4
    assert t["realized_accepted_with_difference"] == 1


def test_realized_totals_cover_only_orders_with_an_accepted_invoice(client):  # noqa: F811
    c = client
    a, b = _po(c), _po(c)
    _base(c, a)
    _base(c, b, amount_eur=60000)
    _inv(c, a, total_net=45000)
    _accept(c, a)
    t = c.get("/savings").json()["totals"]
    assert t["orders"] == 2 and t["baseline_eur"] == 110000 and t["negotiated_saving_eur"] == 20000
    assert (
        t["realized_orders"] == 1 and t["realized_baseline_eur"] == 50000 and t["realized_saving_eur"] == 5000
    )


def test_the_data_can_offer_a_baseline_and_the_amount_comes_from_the_data(client):  # noqa: F811
    c = client
    po = _po(c)
    sids = [s["supplier_id"] for s in c.get(f"/requests/{po['request_id']}/shortlist").json()]
    _offer(c, po["request_id"], sids[1], [12, 25])  # 6,000 + 50,000 = 56,000
    cand = {x["basis"]: x for x in c.get(f"/orders/{po['id']}/savings").json()["candidates"]}
    assert cand["highest_offer"]["amount_eur"] == 56000
    r = c.put(
        f"/orders/{po['id']}/savings/baseline",
        json={"set_by": "kartel", "basis": "highest_offer", "amount_eur": 1},  # an amount is ignored
    )
    assert r.status_code == 200 and r.json()["baseline"]["amount_eur"] == 56000
    # a later, dearer offer does not rewrite the saved baseline
    _offer(c, po["request_id"], sids[2], [20, 30])
    assert c.get(f"/orders/{po['id']}/savings").json()["baseline"]["amount_eur"] == 56000


def test_a_basis_the_data_cannot_give_is_refused_and_a_typed_one_needs_a_note(client):  # noqa: F811
    c = client
    po = _po(c)  # one offer only: no "dearest offer" to point at
    assert _base(c, po, basis="highest_offer").status_code == 409
    assert _base(c, po, note="").status_code == 422
    assert _base(c, po, amount_eur=0).status_code == 422
    assert _base(c, po, basis="guess").status_code == 422
    assert _base(c, po, set_by=" ").status_code == 409
    assert c.get("/orders/999/savings").status_code == 404


def test_changing_or_clearing_a_baseline_is_logged(client):  # noqa: F811
    c = client
    po = _po(c)
    _base(c, po)
    _base(c, po, amount_eur=52000, note="new list")
    assert c.get(f"/orders/{po['id']}/savings").json()["baseline"]["amount_eur"] == 52000
    r = c.delete(f"/orders/{po['id']}/savings/baseline", params={"cleared_by": "kartel"})
    assert r.status_code == 200 and r.json()["baseline"] is None
    assert c.get("/savings").json()["rows"] == []
    actions = [
        a.action for a in c.session.query(AuditLog).order_by(AuditLog.id) if a.action.startswith("savings.")
    ]
    assert actions == ["savings.baseline", "savings.baseline", "savings.baseline_cleared"]
    second = c.session.query(AuditLog).filter_by(action="savings.baseline").all()[1]
    assert second.payload["before"] == {"basis": "manual", "amount_eur": 50000}


def test_the_excel_export_carries_the_rows_and_the_totals(client):  # noqa: F811
    c = client
    a = _po(c)
    _base(c, a)
    _inv(c, a, total_net=45000)
    _accept(c, a)
    r = c.get("/savings/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:2] == ("Order", "Supplier") and rows[1][0] == f"PO #{a['id']}"
    assert rows[1][5] == 50000 and rows[1][7] == 5000 and rows[1][12] == 5000
    assert any(r[0] == "Realized" and r[12] == 5000 for r in rows)
