"""The legal module: norm base and loader (scripted sites, no internet), Incoterms, the playbook on two invented
contracts, decisions, the optional AI read with a fake model, reports, the portfolio and the sign-off gate."""

# ruff: noqa: E501, F811
import base64
import io
import zipfile
from datetime import date, timedelta
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from app.ai.base import AIResult
from app.ai.tasks import set_adapter_factory
from app.config import get_settings
from app.legal import api as legal_api
from app.legal import incoterms, loader, playbook
from app.legal.catalog import CATALOG
from app.main import app
from tests.test_pipeline_e2e import SETTINGS, client, new_request, run_to_offers  # noqa: F401

FIX = Path(__file__).parent / "fixtures" / "legal"
BUY = (FIX / "supply_buy.txt").read_text()
SELL = (FIX / "sale_sell.txt").read_text()

# An invented act in the shape of an EUR-Lex page: a table of contents that repeats the numbers, recitals,
# articles with the title on the next line, a sentence that merely mentions an article, and an article 12g.
ACT_HTML = """<html><head><title>x</title><style>p{}</style></head><body>
<p>THE EXAMPLE DIRECTIVE ON TESTING contains recitals that explain why the invented rules below exist and who
they bind, which is only here so the preamble has some length to be kept by the parser as one block.</p>
<p>TABLE OF CONTENTS</p><p>Article 1</p><p>Article 2</p><p>Article 3</p>
<p class="oj-ti-art">Article 1</p><p class="oj-sti-art">Subject matter</p><p>This invented act sets rules for tests.</p>
<p class="oj-ti-art">Article 2</p><p class="oj-sti-art">Scope</p><p>It applies to every test. Article 3 of this act sets the term.</p>
<p>Article 3 of this act sets the term for payment.</p>
<p class="oj-ti-art">Article 3</p><p class="oj-sti-art">Payment term</p><table><tr><td>1.</td><td>The term shall not exceed sixty days.</td></tr></table>
<p class="oj-ti-art">Article 12g</p><p class="oj-sti-art">Contract clause</p><p>The seller shall include a clause in the contract.</p>
</body></html>"""


def b64(data: bytes | str) -> str:
    return base64.b64encode(data if isinstance(data, bytes) else data.encode()).decode()


def add_contract(c, text=BUY, role="buy", **over):
    body = {
        "by": "kartel",
        "title": "Test contract",
        "counterparty": "Brightwave",
        "role": role,
        "text": text,
        **over,
    }
    r = c.post("/legal/contracts", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def keys(contract):
    return [f["rule_key"] for f in contract["findings"]]


# ---- the base ------------------------------------------------------------------------------------


def test_catalogue_covers_the_five_jurisdictions_and_starts_empty(client):
    s = client.get("/legal/sources").json()
    assert s["total"] == len(CATALOG) >= 50 and s["loaded"] == 0 and s["articles"] == 0
    assert {x["jurisdiction"] for x in s["sources"]} == {"EU", "INT", "RS", "CN", "AE"}
    assert next(x for x in s["sources"] if x["key"] == "cisg")["status"] == "not_loaded"
    assert next(x for x in s["sources"] if x["key"] == "eu-2011-7")["ref"] == "32011L0007"
    assert all(x["check_name"] for x in s["sources"] if x["jurisdiction"] in ("RS", "CN", "AE"))
    assert len({e["key"] for e in CATALOG}) == len(CATALOG)


def test_parser_keeps_real_articles_and_drops_table_of_contents_and_mentions():
    parsed = loader.parse_provisions(loader.html_to_text(ACT_HTML))
    assert [p.number for p in parsed] == ["preamble", "1", "2", "3", "12g"]
    by = {p.number: p for p in parsed}
    assert by["1"].heading == "Subject matter" and "invented act" in by["1"].text
    assert by["3"].heading == "Payment term" and "sixty days" in by["3"].text
    assert "Article 3 of this act sets the term." in by["2"].text  # a mention stays inside its article
    assert by["12g"].heading == "Contract clause"
    assert loader.parse_provisions("just one block of text")[0].number == "text"


def scripted(handler):
    return httpx.MockTransport(handler)


def test_fetch_from_eurlex_stores_articles_and_audits(client, monkeypatch):
    seen = []

    def handler(req: httpx.Request):
        seen.append(str(req.url))
        return httpx.Response(200, content=ACT_HTML.encode(), headers={"content-type": "text/html"})

    monkeypatch.setattr(legal_api, "TRANSPORT", scripted(handler))
    r = client.post("/legal/sources/eu-2011-7/fetch", json={"by": "kartel"})
    assert r.status_code == 200 and r.json()["articles"] == 5
    assert any(u.startswith("https://publications.europa.eu/resource/celex/32011L0007") for u in seen)
    src = client.get("/legal/sources/eu-2011-7").json()
    assert src["status"] == "loaded" and src["loaded_by"] == "kartel" and len(src["text_hash"]) == 64
    assert [p["number"] for p in src["provisions"]][-1] == "12g"
    assert client.post("/legal/sources/eu-2011-7/fetch", json={"by": " "}).status_code == 409


def test_fetch_prefers_the_newest_consolidated_text(client, monkeypatch):
    seen = []
    versions = ["02011L0007-20110216", "02011L0007-20130101", "02011L0007-29991231", "02011L0007"]

    def handler(req: httpx.Request):
        seen.append(str(req.url))
        if "sparql" in req.url.path:
            bindings = [{"c": {"value": v}} for v in versions]
            return httpx.Response(200, json={"results": {"bindings": bindings}})
        return httpx.Response(200, content=ACT_HTML.encode(), headers={"content-type": "text/html"})

    monkeypatch.setattr(legal_api, "TRANSPORT", scripted(handler))
    assert client.post("/legal/sources/eu-2011-7/fetch", json={"by": "kartel"}).status_code == 200
    assert seen[-1].endswith("/celex/02011L0007-20130101")  # not the future date, not the original
    assert client.get("/legal/sources/eu-2011-7").json()["note"] == "Consolidated text 20130101"


def test_fetch_explains_browser_checks_and_blocked_networks(client, monkeypatch):
    monkeypatch.setattr(
        legal_api,
        "TRANSPORT",
        scripted(lambda req: httpx.Response(202, text="<html>JavaScript is required</html>")),
    )
    r = client.post("/legal/sources/eu-2011-7/fetch", json={"by": "kartel"})
    assert r.status_code == 409 and "browser check" in r.json()["detail"] and "upload" in r.json()["detail"]
    monkeypatch.setattr(legal_api, "TRANSPORT", scripted(lambda req: httpx.Response(403)))
    r = client.post("/legal/sources/eu-2011-7/fetch", json={"by": "kartel"})
    assert r.status_code == 409 and "network may block" in r.json()["detail"]
    r = client.post("/legal/sources/rs-zoo/fetch", json={"by": "kartel"})
    assert r.status_code == 409 and "no download address" in r.json()["detail"]
    assert client.post("/legal/sources/nope/fetch", json={"by": "kartel"}).status_code == 404
    assert client.get("/legal/sources").json()["loaded"] == 0


def test_upload_search_and_citation_states(client):
    r = client.post(
        "/legal/sources/eu-2011-7/upload",
        json={"filename": "dir.html", "content_base64": b64(ACT_HTML), "by": "kartel"},
    )
    assert r.status_code == 200 and r.json()["articles"] == 5
    hits = client.get("/legal/search", params={"q": "sixty days"}).json()
    assert hits and hits[0]["source"] == "eu-2011-7" and hits[0]["number"] == "3"
    assert client.get("/legal/search", params={"q": "ab"}).status_code == 409
    rules = {r["key"]: r for r in client.get("/legal/rules").json()}
    states = {(c["source"], c["article"]): c["status"] for c in rules["payment_long"]["citations_checked"]}
    assert states[("eu-2011-7", "3")] == "verified" and states[("rs-payment-terms", None)] == "not_loaded"
    late = {
        (c["source"], c["article"]): c["status"] for c in rules["late_interest_missing"]["citations_checked"]
    }
    assert late[("eu-2011-7", "6")] == "article_missing"  # the invented act has no article 6
    assert (
        client.post(
            "/legal/sources/cisg/upload",
            json={"filename": "x.pdf", "content_base64": b64(b"%PDF-1.4 broken"), "by": "kartel"},
        ).status_code
        == 409
    )
    assert (
        client.post(
            "/legal/sources/cisg/upload", json={"filename": "x.txt", "content_base64": "@@@", "by": "kartel"}
        ).status_code
        == 422
    )


def test_loader_command_line_reads_a_folder(client, tmp_path, monkeypatch):
    (tmp_path / "eu-2011-7.html").write_text(ACT_HTML)
    (tmp_path / "unknown-act.txt").write_text("ignored")
    monkeypatch.setattr("app.db.get_session", lambda: iter([client.session]))
    assert loader.main(["--dir", str(tmp_path), "--by", "ops"]) == 0
    assert client.get("/legal/sources/eu-2011-7").json()["loaded_by"] == "ops"


# ---- Incoterms -----------------------------------------------------------------------------------


def test_incoterms_cover_eleven_rules_and_ten_topics(client):
    data = client.get("/legal/incoterms").json()
    assert (
        data["version"] == "2020"
        and [r["code"] for r in data["rules"]] == incoterms.CODES
        and len(incoterms.CODES) == 11
    )
    assert {g["group"] for g in data["rules"]} == {"E", "F", "C", "D"}
    for r in data["rules"]:
        assert [t["key"] for t in r["topics"]] == [k for k, _ in incoterms.TOPICS]
        assert all(t["seller"].strip() and t["buyer"].strip() for t in r["topics"])
        assert r["watch"]["buy"] and r["watch"]["sell"] and "own summary" in r["source"].lower()
    by = {r["code"]: r for r in data["rules"]}
    topic = lambda code, key: next(t for t in by[code]["topics"] if t["key"] == key)  # noqa: E731
    assert "(C)" in topic("CIF", "insurance")["seller"] and "(A)" in topic("CIP", "insurance")["seller"]
    assert (
        "import" in topic("DDP", "clearance")["seller"].lower()
        and "import" in topic("DAP", "clearance")["buyer"].lower()
    )
    assert [
        c for c, r in by.items() if "unload" in topic(c, "delivery")["seller"].lower() and r["code"] == "DPU"
    ] == ["DPU"]
    assert by["FOB"]["sea_only"] and not by["FCA"]["sea_only"]
    assert any("ownership" in n for n in data["not_covered"])


# ---- the playbook --------------------------------------------------------------------------------


def test_every_check_has_a_rule_and_every_reference_exists_in_the_catalogue():
    assert set(playbook.CHECKS) == {r.key for r in playbook.RULES}
    known = {e["key"] for e in CATALOG}
    for r in playbook.RULES:
        assert r.role in playbook.ROLES and r.severity in playbook.SEVERITIES and r.risk and r.suggest
        assert all(c["source"] in known for c in r.citations), r.key


def test_supplier_contract_finds_the_planted_problems_and_quotes_the_text(client):
    c = add_contract(client)
    found = set(keys(c))
    planted = {
        "advance_payment",
        "payment_short",
        "penalty_rate_high",
        "penalty_uncapped",
        "penalty_adjustable",
        "liability_supplier_cap_low",
        "liability_consequential",
        "liability_gross_negligence",
        "warranty_short",
        "claims_period_short",
        "deemed_acceptance_short",
        "auto_renewal",
        "unilateral_change",
        "termination_one_sided",
        "assignment_one_sided",
        "governing_law_outside",
        "cisg_status",
        "incoterms_version",
        "incoterms_mode",
        "incoterms_role_watch",
        "delivery_date_vague",
        "data_processing_missing",
        "data_transfer",
        "sanctions_clause_missing",
        "embargo_mention",
    }
    assert planted <= found, planted - found
    assert (
        not {"payment_long", "warranty_long", "liability_cap_missing", "retention_of_title_missing"} & found
    )  # sell-side rules
    assert all(f["quote"] in c["text"] for f in c["findings"] if f["quote"])
    sev = {f["rule_key"]: f["severity"] for f in c["findings"]}
    assert (
        sev["advance_payment"] == "red"
        and sev["embargo_mention"] == "red"
        and sev["penalty_rate_high"] == "red"
    )
    assert sev["unilateral_change"] == "red"  # "without prior notice"
    assert [f["severity"] for f in c["findings"]] == sorted(
        (f["severity"] for f in c["findings"]), key=["red", "yellow", "info"].index
    )
    red, yellow = c["counts"]["red"], c["counts"]["yellow"]
    assert c["score"] == 10 * red + 3 * yellow and c["status"] == "analysed"
    assert (
        c["summary"]["law"] == "People's Republic of China"
        and c["summary"]["cisg_excluded"]
        and c["summary"]["incoterms"] == ["FOB Shanghai"]
    )
    assert c["summary"]["forum"] == "arbitration: Shenzhen" and c["summary"]["warranty_months"] == 6
    china = next(f for f in c["findings"] if f["rule_key"] == "penalty_adjustable")
    assert (
        china["citations"][0]["source"] == "cn-civil-code" and china["citations"][0]["status"] == "not_loaded"
    )
    watch = next(f for f in c["findings"] if f["rule_key"] == "incoterms_role_watch")
    assert "container" in watch["risk"].lower() or "vessel" in watch["risk"].lower()


def test_customer_contract_uses_the_sell_side_rules(client):
    c = add_contract(client, SELL, "sell", counterparty="Gulf Horizon")
    found = {f["rule_key"]: f for f in c["findings"]}
    assert {
        "payment_long",
        "late_interest_missing",
        "liability_cap_missing",
        "warranty_long",
        "retention_of_title_missing",
        "incoterms_conflict",
        "jurisdiction_missing",
    } <= set(found)
    assert found["liability_cap_missing"]["severity"] == "red"
    assert (
        found["payment_long"]["title"].endswith("90 days") and found["payment_long"]["severity"] == "yellow"
    )
    assert (
        found["incoterms_role_watch"]["severity"] == "yellow"
        and "import" in found["incoterms_role_watch"]["risk"].lower()
    )
    assert not {"advance_payment", "payment_short", "warranty_short", "unilateral_change"} & set(found)
    dat = [f for f in c["findings"] if f["rule_key"] == "incoterms_version" and f["severity"] == "red"]
    assert dat and "DPU" in dat[0]["risk"]
    assert c["summary"]["payment_days"] == 90 and c["summary"]["law"] == "United Arab Emirates"
    uae = next(f for f in c["findings"] if f["rule_key"] == "penalty_adjustable")
    assert uae["citations"][0]["source"] == "ae-civil"


def test_incoterms_without_place_or_edition(client):
    text = "1. Delivery\n1.1 The goods are delivered CIF to the buyer. The goods travel by air freight.\n2. Payment\n2.1 Payment is due within 30 days."
    c = add_contract(client, text, "buy")
    found = set(keys(c))
    assert {"incoterms_place_missing", "incoterms_version", "incoterms_mode"} <= found
    c2 = add_contract(
        client,
        "1. Delivery\n1.1 The goods are delivered FCA Belgrade warehouse (Incoterms 2020). Payment is due within 45 days of invoice.",
        "buy",
    )
    assert not {"incoterms_place_missing", "incoterms_version", "incoterms_mode", "incoterms_missing"} & set(
        keys(c2)
    )
    c3 = add_contract(
        client,
        "1. Goods\n1.1 The supplier will send the goods to our warehouse in Belgrade within two weeks of the order.",
        "buy",
    )
    assert "incoterms_missing" in keys(c3)


# ---- decisions, rules, settings ------------------------------------------------------------------


def test_decisions_need_a_name_lower_the_score_and_survive_a_new_check(client):
    c = add_contract(client)
    cid = c["id"]
    red = next(f for f in c["findings"] if f["rule_key"] == "advance_payment")
    r = client.post(
        f"/legal/findings/{red['id']}/decision",
        json={"state": "dismissed", "note": "prepaid by agreement", "by": " "},
    )
    assert r.status_code == 409
    assert (
        client.post(
            f"/legal/findings/{red['id']}/decision", json={"state": "maybe", "by": "kartel"}
        ).status_code
        == 409
    )
    d = client.post(
        f"/legal/findings/{red['id']}/decision",
        json={"state": "dismissed", "note": "prepaid by agreement", "by": "kartel"},
    ).json()
    assert d["state"] == "dismissed" and d["decided_by"] == "kartel" and d["decided_at"]
    after = client.get(f"/legal/contracts/{cid}").json()
    assert after["score"] == c["score"] - 10 and after["counts"]["red"] == c["counts"]["red"] - 1
    again = client.post(f"/legal/contracts/{cid}/analyse", json={"by": "kartel"}).json()
    kept = next(f for f in again["findings"] if f["rule_key"] == "advance_payment")
    assert (
        kept["state"] == "dismissed"
        and kept["note"] == "prepaid by agreement"
        and again["score"] == after["score"]
    )
    lawyer = next(f for f in again["findings"] if f["rule_key"] == "liability_consequential")
    client.post(f"/legal/findings/{lawyer['id']}/decision", json={"state": "to_lawyer", "by": "kartel"})
    assert (
        client.get(f"/legal/contracts/{cid}").json()["score"] == after["score"]
    )  # still open for the lawyer


def test_rule_thresholds_and_switches_change_the_result(client):
    first = add_contract(client)
    assert "warranty_short" in keys(first)
    r = client.patch(
        "/legal/rules/warranty_short",
        json={"changed_by": "kartel", "values": {"params": {"min_months": 3}, "draft": False}},
    )
    assert r.status_code == 200 and r.json()["params"]["min_months"] == 3 and r.json()["draft"] is False
    assert "warranty_short" not in keys(add_contract(client))
    client.patch("/legal/rules/embargo_mention", json={"changed_by": "kartel", "values": {"enabled": False}})
    assert "embargo_mention" not in keys(add_contract(client))
    bad = client.patch(
        "/legal/rules/warranty_short", json={"changed_by": "kartel", "values": {"params": {"nope": 1}}}
    )
    assert bad.status_code == 409
    assert (
        client.patch(
            "/legal/rules/warranty_short", json={"changed_by": "kartel", "values": {"severity": "purple"}}
        ).status_code
        == 409
    )
    assert (
        client.patch(
            "/legal/rules/warranty_short", json={"changed_by": "", "values": {"enabled": True}}
        ).status_code
        == 409
    )


def test_settings_approved_laws_weights_and_signoff(client):
    base = client.get("/legal/settings").json()
    assert (
        "serbia" in base["approved_laws"]
        and base["weights"] == {"red": 10, "yellow": 3, "info": 0}
        and base["require_signoff"] is False
    )
    bad = client.put("/legal/settings", json={"changed_by": "kartel", "values": {"weights": {"red": 1}}})
    assert bad.status_code == 409
    client.put(
        "/legal/settings",
        json={
            "changed_by": "kartel",
            "values": {"approved_laws": ["China", "Serbia"], "weights": {"red": 20, "yellow": 5, "info": 1}},
        },
    )
    c = add_contract(client)
    assert "governing_law_outside" not in keys(c)
    assert c["score"] == 20 * c["counts"]["red"] + 5 * c["counts"]["yellow"] + 1 * c["counts"]["info"]


def test_contract_input_validation_and_word_files(client):
    base = {"by": "kartel", "title": "T", "role": "buy"}
    assert client.post("/legal/contracts", json={**base, "text": "too short"}).status_code == 409
    assert client.post("/legal/contracts", json={**base, "role": "both", "text": BUY}).status_code == 409
    assert client.post("/legal/contracts", json={**base, "text": BUY, "by": ""}).status_code == 409
    assert client.post("/legal/contracts", json={**base, "text": BUY, "request_id": 999}).status_code == 404
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        paras = "".join(f"<w:p><w:r><w:t>{ln}</w:t></w:r></w:p>" for ln in BUY.splitlines() if ln.strip())
        z.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{paras}</w:body></w:document>',
        )
    r = client.post(
        "/legal/contracts", json={**base, "filename": "supply.docx", "content_base64": b64(buf.getvalue())}
    )
    assert r.status_code == 200 and "advance_payment" in keys(r.json())
    assert client.delete(f"/legal/contracts/{r.json()['id']}", params={"by": "kartel"}).status_code == 200
    assert client.get(f"/legal/contracts/{r.json()['id']}").status_code == 404


# ---- the optional AI read -------------------------------------------------------------------------


class FakeReview:
    def __init__(self, answer):
        self.answer, self.calls = answer, []

    def extract(self, prompt, variables, schema):
        self.calls.append((prompt.id, variables))
        return AIResult(
            parsed=schema.model_validate(self.answer),
            model="fake-model",
            prompt_id=prompt.id,
            prompt_version=prompt.version,
        )


@pytest.fixture
def ai(client):  # noqa: F811
    fake = FakeReview(
        {
            "confidence": 0.9,
            "findings": [
                {
                    "title": "Exclusivity",
                    "quote": "The Buyer may resell the goods in Serbia, Kazakhstan and Russia.",
                    "severity": "yellow",
                    "risk": "Territory is wide.",
                    "suggest": "Limit the territory.",
                },
                {
                    "title": "Invented",
                    "quote": "The Buyer must give the Supplier a kidney.",
                    "severity": "red",
                    "risk": "x",
                    "suggest": "y",
                },
                {"title": "Too short", "quote": "Data", "severity": "info", "risk": "x", "suggest": "y"},
            ],
        }
    )
    cfg = SETTINGS.model_copy(update={"ai_provider": "anthropic", "anthropic_api_key": "test-key"})
    app.dependency_overrides[get_settings] = lambda: cfg
    set_adapter_factory(lambda s: fake)
    yield client, fake
    set_adapter_factory(None)


def test_ai_read_is_opt_in_per_contract_and_only_keeps_real_quotes(ai):
    c, fake = ai
    contract = add_contract(c)
    cid = contract["id"]
    r = c.post(f"/legal/contracts/{cid}/ai-review", json={"by": "kartel"})
    assert r.status_code == 409 and "switched off" in r.json()["detail"] and fake.calls == []
    c.patch(f"/legal/contracts/{cid}", json={"by": "kartel", "values": {"ai_allowed": True}})
    pre = c.get(f"/legal/contracts/{cid}/ai-preview").json()
    assert pre["enabled"] and pre["allowed"] and pre["chars_sent"] == len(contract["text"]) and not pre["cut"]
    r = c.post(f"/legal/contracts/{cid}/ai-review", json={"by": "kartel"}).json()
    assert r["added"] == 1 and r["dropped"] == 2
    ai_f = [f for f in r["contract"]["findings"] if f["by"] == "ai"]
    assert len(ai_f) == 1 and ai_f[0]["quote"] in contract["text"] and ai_f[0]["clause_id"]
    assert fake.calls[0][0] == "contract_review" and "Payment in advance" in fake.calls[0][1]["checked"]
    assert r["contract"]["ai_model"] == "fake-model" and r["contract"]["score"] == contract["score"] + 3
    assert "ai" not in [
        f["rule_key"]
        for f in c.post(f"/legal/contracts/{cid}/analyse", json={"by": "kartel"}).json()["findings"]
        if f["by"] == "rules"
    ]


def test_ai_off_means_a_plain_refusal(client):
    c = add_contract(client, ai_allowed=True)
    r = client.post(f"/legal/contracts/{c['id']}/ai-review", json={"by": "kartel"})
    assert r.status_code == 409 and "not set up" in r.json()["detail"]
    assert client.get(f"/legal/contracts/{c['id']}/ai-preview").json()["enabled"] is False


# ---- reports and portfolio ------------------------------------------------------------------------


def test_reports_list_every_remark_with_risk_advice_and_norms(client):
    c = add_contract(client)
    md = client.get(f"/legal/contracts/{c['id']}/report.md")
    assert md.status_code == 200 and "attachment" in md.headers["content-disposition"]
    assert "[RED] Payment in advance" in md.text and "Ask for:" in md.text and "not legal advice" in md.text
    assert md.text.count("## [") == len(c["findings"])
    x = client.get(f"/legal/contracts/{c['id']}/report.xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    assert [cell.value for cell in ws[1]][:3] == ["Severity", "Remark", "Clause says"] and ws.max_row == len(
        c["findings"]
    ) + 1
    assert client.get("/legal/contracts/999/report.md").status_code == 404


def test_portfolio_finds_bottlenecks_expiries_and_weak_counterparties(client):
    soon = (date.today() + timedelta(days=30)).isoformat()
    add_contract(client, BUY, "buy", counterparty="Brightwave", expires_on=soon)
    add_contract(client, SELL, "sell", counterparty="Gulf Horizon")
    p = client.get("/legal/portfolio").json()
    assert p["contracts"] == 2 and p["analysed"] == 2 and p["open_red"] >= 4
    assert p["average_payment_days"] == 52.5  # 15 and 90 days
    assert [e["title"] for e in p["expiring_90_days"]] == ["Test contract"] and len(p["auto_renewing"]) == 2
    assert p["bottlenecks"][0]["contracts"] >= 1 and {t["counterparty"] for t in p["counterparties"]} == {
        "Brightwave",
        "Gulf Horizon",
    }
    assert {row["law"] for row in p["table"]} == {"People's Republic of China", "United Arab Emirates"}
    assert client.get("/legal/contracts").json()[0]["counts"]["open"] > 0


# ---- the sign-off gate ----------------------------------------------------------------------------


def accepted_quote(c):
    rid = new_request(c)
    run_to_offers(c, rid)
    best = c.post(f"/requests/{rid}/score").json()[0]
    q = c.post(
        f"/requests/{rid}/quotes",
        json={"offer_id": best["offer_id"], "markup": 0.2, "currency": "EUR", "created_by": "kartel"},
    ).json()
    c.post(f"/quotes/{q['id']}/decision", json={"approve": True, "decided_by": "kartel"})
    assert (
        c.post(f"/quotes/{q['id']}/outcome", json={"kind": "accepted", "confirmed_by": "kartel"}).status_code
        == 200
    )
    return rid


def test_signoff_rule_holds_the_order_until_red_remarks_are_dealt_with(client):
    c = client
    rid = accepted_quote(c)
    contract = add_contract(c, request_id=rid)
    summary = c.get(f"/requests/{rid}/legal").json()
    assert summary["open_red"] == contract["counts"]["red"] > 0 and summary["require_signoff"] is False
    c.put("/legal/settings", json={"changed_by": "kartel", "values": {"require_signoff": True}})
    sent = len(c.outbox)
    r = c.post(f"/requests/{rid}/orders", json={"approved_by": "kartel"})
    assert (
        r.status_code == 409 and "legal sign-off is required" in r.json()["detail"] and len(c.outbox) == sent
    )
    for f in contract["findings"]:
        if f["severity"] == "red":
            c.post(
                f"/legal/findings/{f['id']}/decision",
                json={"state": "accepted", "note": "known", "by": "kartel"},
            )
    assert c.get(f"/requests/{rid}/legal").json()["open_red"] == 0
    assert c.post(f"/requests/{rid}/orders", json={"approved_by": "kartel"}).status_code == 200


def test_without_the_rule_a_red_remark_never_stops_an_order(client):
    rid = accepted_quote(client)
    add_contract(client, request_id=rid)
    assert client.post(f"/requests/{rid}/orders", json={"approved_by": "kartel"}).status_code == 200


@pytest.mark.integration
def test_contract_lifecycle_against_real_postgres(pg_url):
    """Postgres enforces the foreign keys that SQLite ignores: analyse again, then delete, children first."""
    from sqlalchemy import create_engine, func, select
    from sqlalchemy.orm import Session

    from app.legal import service as svc
    from app.models import Contract, ContractClause, ContractFinding

    with Session(create_engine(pg_url)) as s:
        svc.ensure_seed(s)
        c = svc.create_contract(s, "pg tester", title="pg lifecycle", counterparty="x", role="buy", text=BUY)
        first = svc.analyse(s, "pg tester", c.id)
        assert first["findings"]
        again = svc.analyse(s, "pg tester", c.id)
        assert len(again["findings"]) == len(first["findings"])
        svc.delete_contract(s, "pg tester", c.id)
        assert s.get(Contract, c.id) is None
        for model in (ContractClause, ContractFinding):
            assert s.scalar(select(func.count()).select_from(model).where(model.contract_id == c.id)) == 0


def test_parser_reads_dotted_and_roman_article_numbers_and_ignores_cross_references():
    unidroit = (
        "UNIDROIT Principles 2016\nARTICLE 1.1\n(Freedom of contract)\n"
        "The parties are free to enter into a contract and to determine its content.\n"
        "ARTICLE 1.2\n(No form required)\nNothing in these Principles requires a contract to be in writing.\n"
        "Article 76 CISG, is to facilitate proof of harm where no replacement has been made"
    )
    got = loader.parse_provisions(unidroit)
    assert [p.number for p in got] == ["1.1", "1.2"]
    assert got[0].heading == "(Freedom of contract)" and got[0].text.startswith("The parties are free")
    roman = "Article I\n1. This Convention shall apply to arbitral awards.\nArticle II\n1. Each State shall recognise an agreement in writing."
    assert [p.number for p in loader.parse_provisions(roman)] == ["i", "ii"]


def test_pdf_style_article_without_a_title_keeps_its_first_sentence_in_the_text():
    got = loader.parse_provisions(
        "Article 78\nIf a party fails to pay the price, the other party is entitled to interest.\nArticle 79\nA party is not liable."
    )
    assert got[0].heading == "" and got[0].text.startswith("If a party fails")
